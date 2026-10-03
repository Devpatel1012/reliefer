import io
import logging
import re
from typing import Any

logger = logging.getLogger("reliefer.resume_parser")


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract raw text from PDF file bytes using pdfplumber with pypdf fallback."""
    text_parts = []

    # Primary: pdfplumber
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        if text_parts:
            return "\n".join(text_parts)
    except Exception as e:
        logger.warning(f"pdfplumber extraction failed, trying pypdf fallback: {e}")

    # Fallback: pypdf
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        if text_parts:
            return "\n".join(text_parts)
    except Exception as e:
        logger.error(f"pypdf extraction failed: {e}")

    return ""


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract raw text from DOCX file bytes using python-docx."""
    text_parts = []
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))
        return "\n".join(text_parts)
    except Exception as e:
        logger.error(f"docx extraction failed: {e}")
        return ""


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract text from supported file formats (PDF, DOCX, TXT)."""
    fn_lower = filename.lower()
    if fn_lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif fn_lower.endswith((".docx", ".doc")):
        return extract_text_from_docx(file_bytes)
    else:
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="ignore")


# ─── Section patterns ─────────────────────────────────────────────────────────

SECTION_HEADER_PATTERNS = {
    "skills": re.compile(
        r"^(technical\s+skills?|skills?|technologies|expertise|core\s+competencies|"
        r"tools?\s*(&|and|\+)\s*technologies|programming\s+languages?|competencies|"
        r"languages?\s*&?\s*tools?|tech\s+stack)",
        re.IGNORECASE
    ),
    "experience": re.compile(
        r"^(work\s+experience|professional\s+experience|experience|"
        r"employment(\s+history)?|work\s+history|career(\s+history)?|internships?)",
        re.IGNORECASE
    ),
    "education": re.compile(
        r"^(education(al)?\s*(background|qualifications?)?|academic(\s+background)?|"
        r"qualifications?|degrees?|schooling|training)",
        re.IGNORECASE
    ),
    "projects": re.compile(
        r"^(projects?|personal\s+projects?|key\s+projects?|academic\s+projects?|"
        r"side\s+projects?|portfolio|notable\s+projects?|selected\s+projects?)",
        re.IGNORECASE
    ),
    "achievements": re.compile(
        r"^(achievements?|awards?|honors?|certifications?|licenses?\s*(&|and|\+)?\s*certifications?|"
        r"certifications?\s*(&|and|\+)?\s*licenses?|accomplishments?|"
        r"publications?|research\s+publications?|certifications?\s*(&|and|\+)\s*achievements?|"
        r"awards?\s*(&|and|\+)\s*certifications?)",
        re.IGNORECASE
    ),
    "other": re.compile(
        r"^(summary|professional\s+summary|profile|about(\s+me)?|"
        r"additional\s+info(rmation)?|interests?|extracurriculars?)",
        re.IGNORECASE
    ),
}

KNOWN_TECH_STACK = [
    "Python", "JavaScript", "TypeScript", "Java", "C++", "C#", "Go", "Rust",
    "Swift", "Kotlin", "PHP", "Ruby", "R", "MATLAB", "Scala", "Perl", "Bash",
    "HTML", "CSS", "SQL", "React", "Vue", "Angular", "Next.js", "Node.js",
    "Django", "Flask", "FastAPI", "Spring", "Express", "Laravel", "Rails",
    "PostgreSQL", "MySQL", "MongoDB", "Redis", "SQLite", "Oracle", "Cassandra",
    "Docker", "Kubernetes", "AWS", "GCP", "Azure", "Terraform", "Ansible",
    "Git", "CI/CD", "GitHub Actions", "Jenkins", "Linux",
    "TensorFlow", "PyTorch", "Keras", "Scikit-learn", "OpenCV", "Pandas",
    "NumPy", "Matplotlib", "Seaborn", "LangChain", "Hugging Face", "CUDA",
    "GraphQL", "REST API", "gRPC", "Kafka", "Spark", "Airflow",
    "Streamlit", "Jupyter", "Colab", "MLflow", "n8n", "Vector Database",
    "Transformers", "LLaMA", "CvT", "OpenAI", "Agentic AI", "SVLM"
]

ROLE_KEYWORDS_RE = re.compile(
    r"\b(engineer|developer|analyst|manager|intern|lead|architect|"
    r"scientist|designer|consultant|specialist|director|officer|"
    r"associate|assistant|coordinator|researcher|student|fellow)\b",
    re.IGNORECASE
)

DATE_PATTERN_RE = re.compile(
    r"((?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*\d{4}|\b(19|20)\d{2}\b)"
    r"\s*[\-–—to\s]+\s*"
    r"((?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*\d{4}|\b(19|20)\d{2}\b|present|expected|current)",
    re.IGNORECASE
)


def _clean_text(s: str) -> str:
    s = re.sub(r"\[\s*↗\s*\]|\[\s*Certificate\s*\]", "", s)
    s = re.sub(r"^[•\-\*►▪▸◦◉●]\s*", "", s)
    return s.strip()


def _split_into_sections(lines: list[str]) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {
        "header": [],
        "skills": [],
        "experience": [],
        "education": [],
        "projects": [],
        "achievements": [],
        "other": [],
    }
    current_section = "header"

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        clean_line = re.sub(r"[^\w\s&+\-]", "", stripped).strip()
        matched_section = None

        if 2 <= len(clean_line) <= 60:
            for sec_name, pattern in SECTION_HEADER_PATTERNS.items():
                if pattern.match(clean_line):
                    matched_section = sec_name
                    break

        if matched_section:
            current_section = matched_section
        else:
            sections[current_section].append(stripped)

    return sections


def _parse_skills(sections: dict[str, list[str]], raw_text: str) -> list[str]:
    skills_list = []
    CATEGORY_PREFIX_RE = re.compile(
        r"^(languages?|frameworks?|libraries|tools?|technologies|cloud|databases?|"
        r"platforms?|methodologies|soft\s+skills?|hard\s+skills?|other|"
        r"programming|scripting|ide|devops|ml\s*/\s*ai|ai\s*/?\s*ml|"
        r"web|mobile|os|operating\s+systems?|version\s+control|testing|"
        r"coursework|relevant\s+skills?|tech\s+stack|technical\s+expertise|interests)\s*[:\-]?\s*",
        re.IGNORECASE
    )

    skill_lines = sections.get("skills", [])
    for raw_line in skill_lines:
        line_clean = _clean_text(raw_line)
        line_clean = CATEGORY_PREFIX_RE.sub("", line_clean).strip()
        if not line_clean:
            continue
        raw_skills = re.split(r"[,;|•·/]", line_clean)
        for s in raw_skills:
            s = CATEGORY_PREFIX_RE.sub("", s)
            s_clean = re.sub(r"^[\s\-:•·*]+|[\s\-:•·*]+$", "", s).strip()
            s_clean = re.sub(r"\s*[\(\[].+?[\)\]]", "", s_clean).strip()
            if s_clean and 1 < len(s_clean) < 45:
                skills_list.append(s_clean)

    seen = set()
    deduped = []
    for s in skills_list:
        key = s.lower()
        if key not in seen:
            seen.add(key)
            deduped.append(s)

    text_lower = raw_text.lower()
    for tech in KNOWN_TECH_STACK:
        if tech.lower() not in seen:
            if re.search(r"\b" + re.escape(tech.lower()) + r"\b", text_lower):
                seen.add(tech.lower())
                deduped.append(tech)

    return deduped[:30]


def _parse_experience(exp_lines: list[str]) -> list[dict[str, Any]]:
    blocks = []
    curr = []
    for line in exp_lines:
        line_s = line.strip()
        if not line_s:
            continue
        is_bullet = bool(re.match(r"^[•\-\*►▪▸◦◉●]\s*", line_s))
        has_date = bool(re.search(r"\b(19|20)\d{2}\b|(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s*\d{4}|Present|Current", line_s, re.IGNORECASE))
        has_role = bool(ROLE_KEYWORDS_RE.search(line_s))

        if not is_bullet and (has_date or has_role) and len(curr) > 0:
            has_bullets_in_curr = any(re.match(r"^[•\-\*►▪▸◦◉●]\s*", l) for l in curr)
            if has_bullets_in_curr or has_date:
                blocks.append(curr)
                curr = []
        curr.append(line_s)
    if curr:
        blocks.append(curr)

    experience_list = []
    for block in blocks:
        headers = []
        bullets = []
        for l in block:
            if re.match(r"^[•\-\*►▪▸◦◉●]\s*", l):
                bullets.append(re.sub(r"^[•\-\*►▪▸◦◉●]\s*", "", l).strip())
            else:
                if bullets and not re.search(r"\b(19|20)\d{2}\b", l):
                    bullets[-1] += " " + l.strip()
                else:
                    headers.append(l.strip())

        full_h = " ".join(headers)
        date_match = DATE_PATTERN_RE.search(full_h)
        dates = date_match.group(0) if date_match else ""

        clean_headers = []
        for h in headers:
            h_no_date = DATE_PATTERN_RE.sub("", h)
            h_clean = _clean_text(h_no_date)
            h_clean = re.sub(r"\b(Remote|Ahmedabad|Hybrid|On-site)\b", "", h_clean, flags=re.IGNORECASE).strip()
            h_clean = re.sub(r"[\-–—\s]+$", "", h_clean).strip()
            if h_clean:
                clean_headers.append(h_clean)

        role = ""
        company = ""
        if len(clean_headers) >= 2:
            if ROLE_KEYWORDS_RE.search(clean_headers[1]):
                role = clean_headers[1]
                company = clean_headers[0]
            else:
                role = clean_headers[0]
                company = clean_headers[1]
        elif len(clean_headers) == 1:
            h = clean_headers[0]
            parts = re.split(r"\s*[|@—–]\s*|\s+at\s+", h)
            if len(parts) >= 2:
                if ROLE_KEYWORDS_RE.search(parts[1]):
                    role = parts[1]
                    company = parts[0]
                else:
                    role = parts[0]
                    company = parts[1]
            else:
                role = h
                company = "Company"
        else:
            role = "Role"
            company = "Company"

        role = re.sub(r"[\-–—\s]+$", "", role).strip()
        company = re.sub(r"[\-–—\s]+$", "", company).strip()

        desc = " ".join(bullets).strip()

        experience_list.append({
            "role": role[:100],
            "company": company[:100],
            "dates": dates[:60],
            "description": desc[:600],
        })
    return experience_list[:10]


def _parse_education(edu_lines: list[str]) -> list[dict[str, Any]]:
    education_list = []
    current_edu: dict | None = None

    degree_keywords = [
        "bachelor", "master", "phd", "ph.d", "b.s", "m.s", "b.tech", "m.tech",
        "b.e", "m.e", "mba", "degree", "diploma", "b.sc", "m.sc", "undergraduate",
        "postgraduate", "associate"
    ]
    institution_keywords = [
        "university", "college", "institute", "school", "academy",
        "polytechnic", "iit", "nit", "bits", "mit", "stanford"
    ]

    for line in edu_lines:
        line_clean = _clean_text(line)
        if not line_clean:
            continue

        is_degree = any(kw in line_clean.lower() for kw in degree_keywords)
        is_inst = any(kw in line_clean.lower() for kw in institution_keywords)
        
        # Extract GPA / Marks
        gpa_match = re.search(r"(cgpa|gpa|percentage|marks|%)\s*[:\-]?\s*([\d\.]+)", line_clean, re.IGNORECASE)
        gpa = gpa_match.group(2) if gpa_match else ""
        
        # Extract Timing
        date_match = re.search(r"((?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*\d{4}|\b(19|20)\d{2}\b)", line_clean, re.IGNORECASE)
        year = date_match.group(0) if date_match else ""

        if is_degree or is_inst:
            if is_inst and not is_degree and not current_edu:
                inst_name = re.sub(r"((?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*\d{4}|\b(19|20)\d{2}\b).*$", "", line_clean, flags=re.IGNORECASE).strip()
                current_edu = {
                    "degree": "Degree",
                    "institution": inst_name[:120],
                    "year": year,
                    "marks": gpa
                }
            elif is_degree:
                if current_edu and current_edu["degree"] != "Degree":
                    education_list.append(current_edu)
                    current_edu = None

                parts = re.split(r"[,|–\-]\s*", line_clean)
                degree = parts[0].strip()
                degree = re.sub(r"\((cgpa|gpa|percentage|marks|%).*?\)", "", degree, flags=re.IGNORECASE).strip()

                institution = parts[1].strip() if len(parts) > 1 else ""
                institution = re.sub(r"\((cgpa|gpa|percentage|marks|%).*?\)", "", institution, flags=re.IGNORECASE).strip()
                
                if not current_edu:
                    current_edu = {
                        "degree": degree[:120],
                        "institution": institution[:120],
                        "year": year,
                        "marks": gpa
                    }
                else:
                    current_edu["degree"] = degree[:120]
                    if len(parts) > 1:
                        if not current_edu.get("institution") or current_edu["institution"] == "Institution":
                            current_edu["institution"] = institution[:120]
                    if year and not current_edu.get("year"):
                        current_edu["year"] = year
                    if gpa and not current_edu.get("marks"):
                        current_edu["marks"] = gpa
        elif current_edu:
            if not current_edu.get("institution"):
                current_edu["institution"] = line_clean[:120]
            if year and not current_edu.get("year"):
                current_edu["year"] = year
            if gpa and not current_edu.get("marks"):
                current_edu["marks"] = gpa

    if current_edu:
        education_list.append(current_edu)

    return education_list[:6]


def _parse_projects(proj_lines: list[str]) -> list[dict[str, Any]]:
    projects_list = []
    current_proj = None

    date_pattern = r"^(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|[0-9]{4})\b"

    for line in proj_lines:
        line_s = line.strip()
        if not line_s:
            continue

        is_bullet = bool(re.match(r"^[•\-\*►▪▸◦◉●]\s*", line_s))
        clean_l = _clean_text(line_s)

        # Check if line is purely date or date range (e.g., "Mar 2026", "Jan 2026 - Mar 2026", "2024 - 2027")
        if re.match(date_pattern, clean_l, re.IGNORECASE) and len(clean_l) <= 30:
            if current_proj:
                current_proj["description"] = (current_proj["description"] + " (" + clean_l + ")").strip()
                continue

        if ":" in clean_l:
            parts = clean_l.split(":", 1)
            title_part = parts[0].strip()
            desc_part = parts[1].strip()

            if 2 <= len(title_part) <= 60 and not re.search(r"^(tech|stack|built|technologies)\b", title_part, re.IGNORECASE):
                if current_proj:
                    projects_list.append(current_proj)
                current_proj = {
                    "title": title_part[:100],
                    "description": desc_part,
                    "tech_stack": ""
                }
                continue

        if not is_bullet and 2 <= len(clean_l) <= 90:
            if current_proj and (current_proj["description"] or len(current_proj["title"]) > 0):
                projects_list.append(current_proj)
            current_proj = {
                "title": clean_l[:100],
                "description": "",
                "tech_stack": ""
            }
        elif current_proj:
            current_proj["description"] += (" " + clean_l if current_proj["description"] else clean_l)

    if current_proj:
        projects_list.append(current_proj)

    for p in projects_list:
        p["description"] = p["description"].strip()[:500]
        found_techs = [tech for tech in KNOWN_TECH_STACK if re.search(r"\b" + re.escape(tech) + r"\b", p["description"] + " " + p["title"], re.IGNORECASE)]
        if found_techs:
            p["tech_stack"] = ", ".join(list(dict.fromkeys(found_techs))[:5])

    return projects_list[:10]


def _parse_achievements(achieve_lines: list[str]) -> list[dict[str, Any]]:
    achievements = []
    date_pattern = r"^(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|[0-9]{4})\b"

    for line in achieve_lines:
        line_clean = _clean_text(line)
        if not line_clean or len(line_clean) < 3:
            continue
        
        # Check if line is purely date string
        year_m = re.search(r"\b(19|20)\d{2}\b", line_clean)
        year_str = year_m.group(0) if year_m else ""

        if re.match(date_pattern, line_clean, re.IGNORECASE) and len(line_clean) <= 30:
            if achievements:
                if not achievements[-1].get("date") and year_str:
                    achievements[-1]["date"] = year_str
                achievements[-1]["description"] = (achievements[-1]["description"] + " (" + line_clean + ")").strip()
                continue

        if " — " in line_clean or " - " in line_clean:
            parts = re.split(r"\s*[—\-]\s*", line_clean, maxsplit=1)
            title = parts[0].strip()
            desc = parts[1].strip() if len(parts) > 1 else title
        else:
            title = line_clean[:80]
            desc = line_clean
        achievements.append({
            "title": title[:100],
            "description": desc[:400],
            "date": year_str
        })
    return achievements[:10]


def parse_resume_text_to_json(raw_text: str) -> dict[str, Any]:
    """
    Parse raw resume text into structured JSON with sections:
    skills, experience, education, projects, achievements, other_info.
    """
    if not raw_text or not raw_text.strip():
        return {
            "skills": [],
            "experience": [],
            "education": [],
            "projects": [],
            "achievements": [],
            "other_info": [],
            "raw_text": "",
        }

    lines = [line.strip() for line in raw_text.split("\n")]
    sections = _split_into_sections(lines)

    skills = _parse_skills(sections, raw_text)
    experience = _parse_experience(sections.get("experience", []))
    education = _parse_education(sections.get("education", []))
    projects = _parse_projects(sections.get("projects", []))
    achievements = _parse_achievements(sections.get("achievements", []))
    other_info = [_clean_text(l) for l in sections.get("other", []) if len(_clean_text(l)) > 5]

    return {
        "skills": list(dict.fromkeys(skills)),
        "experience": experience,
        "education": education,
        "projects": projects,
        "achievements": achievements,
        "other_info": other_info,
        "raw_text": raw_text[:3000],
    }
